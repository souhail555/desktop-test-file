import io
import json
import os
import sys
import time
from collections import defaultdict, deque
from pathlib import Path
from threading import Lock
from urllib.request import urlopen

PROJECT_ROOT = Path(__file__).resolve().parent
SRC_DIR = PROJECT_ROOT / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from flask import Flask, Response, jsonify, redirect, render_template, request, session, stream_with_context, url_for
from werkzeug.security import check_password_hash

from my_ai_agent.agent import AIAgent
from my_ai_agent.database import ensure_default_user, get_dashboard_stats, get_messages_for_user, get_user, init_db, save_message

MAX_FILE_SIZE = 15 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 100_000
RECOMMENDED_OLLAMA_MODELS = [
    {"name": "llama3.2:3b", "label": "Llama 3.2 · 3B"},
    {"name": "qwen2.5:3b", "label": "Qwen 2.5 · 3B"},
    {"name": "mistral:7b", "label": "Mistral · 7B"},
]


def extract_uploaded_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension == ".txt":
        text = content.decode("utf-8-sig", errors="replace")
    elif extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        table_rows = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        text = "\n".join(paragraphs + table_rows)
    else:
        raise ValueError("الأنواع المدعومة هي PDF وTXT وDOCX فقط.")

    text = text.strip()
    if not text:
        raise ValueError(f"لم أجد نصاً قابلاً للقراءة في الملف {Path(filename).name}.")
    return text[:MAX_EXTRACTED_CHARACTERS]


def create_app() -> Flask:
    app = Flask(__name__)
    app.config["JSON_SORT_KEYS"] = False
    app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024
    app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "dev-secret-key-change-me")
    init_db()
    ensure_default_user()
    agent = AIAgent()
    chat_requests = defaultdict(deque)
    chat_requests_lock = Lock()

    def is_chat_rate_limited(client_ip: str) -> bool:
        now = time.monotonic()
        with chat_requests_lock:
            recent_requests = chat_requests[client_ip]
            while recent_requests and recent_requests[0] <= now - 60:
                recent_requests.popleft()
            if len(recent_requests) >= 10:
                return True
            recent_requests.append(now)
            return False

    @app.errorhandler(400)
    def bad_request(error):
        return jsonify({"error": "Bad request", "message": str(error.description or error)}), 400

    @app.errorhandler(404)
    def not_found(error):
        return jsonify({"error": "Not found", "message": "The requested endpoint does not exist."}), 404

    @app.errorhandler(500)
    def server_error(error):
        return jsonify({"error": "Internal server error", "message": "Something went wrong on the server."}), 500

    def login_required(view_func):
        def wrapper(*args, **kwargs):
            session.setdefault("user_id", get_user("admin")["id"] if get_user("admin") else 1)
            session.setdefault("username", "Guest")
            return view_func(*args, **kwargs)

        wrapper.__name__ = view_func.__name__
        return wrapper

    def get_public_user_id():
        user = get_user("guest")
        if user:
            return user["id"]
        user_id = ensure_default_user()
        if isinstance(user_id, dict):
            return user_id["id"]
        return 1

    @app.route("/")
    def index():
        return redirect(url_for("ai_home"))

    @app.route("/ai")
    def ai_home():
        session.setdefault("user_id", get_public_user_id())
        session.setdefault("username", "Guest")
        stats = get_dashboard_stats()
        return render_template("dashboard.html", username=session.get("username", "Guest"), stats=stats)

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            username = (request.form.get("username") or "Guest").strip() or "Guest"
            session["user_id"] = get_public_user_id()
            session["username"] = username
            return redirect(url_for("ai_home"))

        return render_template("login.html")

    @app.route("/logout")
    def logout():
        session.clear()
        return redirect(url_for("ai_home"))

    @app.route("/dashboard")
    @login_required
    def dashboard():
        session.setdefault("user_id", get_public_user_id())
        session.setdefault("username", "Guest")
        stats = get_dashboard_stats()
        return render_template("dashboard.html", username=session.get("username", "Guest"), stats=stats)

    @app.route("/history")
    @login_required
    def history():
        user_id = session.get("user_id") or get_public_user_id()
        messages = get_messages_for_user(user_id, limit=50)
        return render_template("history.html", username=session.get("username", "Guest"), messages=messages)

    @app.route("/health")
    def health():
        return jsonify({"status": "ok", "service": "my-ai-agent", "version": "1.0.0"})

    @app.route("/api/status")
    def api_status():
        mode = agent.provider if agent.client is not None else "local"
        return jsonify({
            "connected": agent.is_connected(),
            "model": agent.model,
            "mode": mode,
        })

    @app.route("/api/models")
    def api_models():
        if agent.provider == "ollama":
            models = []
            try:
                tags_url = agent.ollama_base_url.removesuffix("/v1") + "/api/tags"
                with urlopen(tags_url, timeout=2) as response:
                    installed = json.loads(response.read()).get("models", [])
                models = [
                    {"name": item.get("name", ""), "installed": True}
                    for item in installed if item.get("name")
                ]
            except Exception:
                pass

            known_names = {model["name"] for model in models}
            models.extend(
                {**model, "installed": False}
                for model in RECOMMENDED_OLLAMA_MODELS
                if model["name"] not in known_names
            )
            return jsonify({"provider": "ollama", "models": models})

        models = [agent.model]
        if agent.provider == "openai":
            models.extend(["gpt-4o-mini", "gpt-4o", "gpt-4.1-mini", "gpt-4.1"])
        unique_models = list(dict.fromkeys(models))
        return jsonify({"provider": agent.provider, "models": [
            {"name": model, "installed": True} for model in unique_models
        ]})

    @app.route("/api/chat", methods=["POST"])
    def chat():
        client_ip = request.headers.get("CF-Connecting-IP", request.remote_addr or "unknown")
        if is_chat_rate_limited(client_ip):
            return jsonify({"error": "تم الوصول للحد المؤقت: 10 رسائل في الدقيقة. انتظر قليلاً ثم حاول مجدداً."}), 429

        session.setdefault("user_id", get_public_user_id())
        session.setdefault("username", "Guest")

        uploads = []
        if request.mimetype == "multipart/form-data":
            payload = request.form
            uploads = request.files.getlist("files")
        else:
            payload = request.get_json(silent=True) or {}
        if not hasattr(payload, "get"):
            return jsonify({"error": "JSON body must be an object."}), 400

        prompt = str(payload.get("prompt", "") or payload.get("message", "")).strip()
        model = str(payload.get("model") or agent.model).strip() or agent.model
        if not prompt and not uploads:
            return jsonify({"error": "اكتب رسالة أو أرفق ملفاً للبدء."}), 400
        if len(uploads) > 5:
            return jsonify({"error": "يمكنك إرفاق خمسة ملفات كحد أقصى."}), 400

        document_context = []
        filenames = []
        try:
            for upload in uploads:
                if not upload.filename:
                    continue
                content = upload.read(MAX_FILE_SIZE + 1)
                if len(content) > MAX_FILE_SIZE:
                    raise ValueError(f"حجم {Path(upload.filename).name} يتجاوز 15MB.")
                extracted = extract_uploaded_text(upload.filename, content)
                filenames.append(Path(upload.filename).name)
                document_context.append(f"\n\n--- محتوى الملف: {Path(upload.filename).name} ---\n{extracted}")
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except Exception:
            return jsonify({"error": "تعذر قراءة الملف. تأكد أنه غير تالف وبصيغة مدعومة."}), 400

        prompt_for_model = prompt + "".join(document_context)
        stored_prompt = prompt or "راجع الملف المرفق"
        if filenames:
            stored_prompt += "\n\nالملفات: " + ", ".join(filenames)

        if request.form.get("stream") == "1" and agent.provider == "ollama" and agent.client is not None:
            def generate_response():
                answer_parts = []
                try:
                    for token in agent.respond_stream(prompt_for_model, model=model):
                        answer_parts.append(token)
                        yield f"data: {json.dumps({'token': token}, ensure_ascii=False)}\n\n"
                    answer = "".join(answer_parts)
                    if not answer:
                        raise RuntimeError("The AI provider returned an empty response. Please try again.")
                    save_message(session["user_id"], stored_prompt, answer, model=model)
                    yield f"data: {json.dumps({'done': True}, ensure_ascii=False)}\n\n"
                except RuntimeError as exc:
                    yield f"data: {json.dumps({'error': str(exc)}, ensure_ascii=False)}\n\n"

            return Response(
                stream_with_context(generate_response()),
                mimetype="text/event-stream",
                headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
            )

        try:
            response = agent.respond(prompt_for_model, model=model)
            save_message(session["user_id"], stored_prompt, response, model=model)
            return jsonify({"status": "success", "answer": response, "model": model,
                            "mode": agent.provider if agent.client is not None else "local",
                            "files": filenames})
        except RuntimeError as exc:
            return jsonify({"error": str(exc)}), 502
        except Exception as exc:
            app.logger.exception("Chat response failed")
            return jsonify({"error": "Failed to generate a response. Please try again."}), 500

    @app.route("/api/tools", methods=["POST"])
    def tool_call():
        payload = request.get_json(silent=True) or {}
        if not isinstance(payload, dict):
            return jsonify({"error": "JSON body must be an object."}), 400

        tool_name = str(payload.get("tool", "")).strip()
        arguments = payload.get("arguments", {}) or {}

        if not tool_name:
            return jsonify({"error": "Tool name is required."}), 400

        try:
            result = agent.execute_tool(tool_name, arguments)
            return jsonify({"status": "success", "tool": tool_name, "result": result})
        except ValueError as exc:
            return jsonify({"error": "Tool execution failed", "message": str(exc)}), 400
        except Exception as exc:
            return jsonify({"error": "Tool execution error", "message": str(exc)}), 500

    @app.route("/api/review", methods=["POST"])
    def review_code():
        payload = request.get_json(silent=True) or {}
        if not isinstance(payload, dict):
            return jsonify({"error": "JSON body must be an object."}), 400

        code = str(payload.get("code", "")).strip()
        if not code:
            return jsonify({"error": "Code is required."}), 400

        try:
            review = agent.review_code(code)
            return jsonify({"status": "success", "review": review})
        except Exception as exc:
            return jsonify({"error": "Review failed", "message": str(exc)}), 500

    return app


app = create_app()


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.getenv("PORT", "5055")), debug=False, use_reloader=False)
